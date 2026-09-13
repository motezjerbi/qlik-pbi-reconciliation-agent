"""
Générateur de rapport d'audit Word (.docx) — charte Talan.

Produit un livrable client complet : page de garde, résumé exécutif,
score de santé + graphique de répartition des constats par criticité,
et détail structuré de tous les constats (origine, détail, diagnostic,
recommandation).

Dépendances : python-docx, matplotlib (déjà utilisées ailleurs dans
l'écosystème Python, à ajouter dans requirements.txt si absentes :
`pip install python-docx matplotlib`).
"""
import io
import datetime
from pathlib import Path
from typing import Dict, List, Optional

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ============================================================
# Charte Talan — alignée sur les variables CSS de app.py
# ============================================================
COLOR_INK = RGBColor(0x0B, 0x12, 0x20)
COLOR_INK_SOFT = RGBColor(0x64, 0x74, 0x8B)
COLOR_ACCENT = RGBColor(0x1E, 0x3A, 0x5F)
COLOR_ACCENT2 = RGBColor(0x25, 0x63, 0xEB)
COLOR_BRASS = RGBColor(0xB4, 0x53, 0x09)
COLOR_DANGER = RGBColor(0xDC, 0x26, 0x26)
COLOR_WARN = RGBColor(0xD9, 0x77, 0x06)
COLOR_OK = RGBColor(0x05, 0x96, 0x69)
COLOR_NEUTRAL = RGBColor(0x64, 0x74, 0x8B)

CRIT_STYLE = {
    "BLOQUANT": (COLOR_DANGER, "FEE2E2"),
    "MAJEUR": (COLOR_WARN, "FEF3C7"),
    "MINEUR": (COLOR_NEUTRAL, "F1F5F9"),
    "OK": (COLOR_OK, "D1FAE5"),
}

NIVEAU_COLORS = {
    "CRITIQUE": COLOR_DANGER,
    "ÉLEVÉ": COLOR_WARN,
    "MODÉRÉ": COLOR_WARN,
    "FAIBLE": COLOR_OK,
    "AUCUN": COLOR_OK,
}


# ============================================================
# HELPERS BAS NIVEAU (OOXML)
# ============================================================
def _shade_cell(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_cell_text(cell, text, color: Optional[RGBColor] = None, bold=False, size=10, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text) if text is not None else "")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Calibri"
    if color is not None:
        run.font.color.rgb = color


def _add_top_border(paragraph, color_hex: str, size: int = 18):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), str(size))
    top.set(qn("w:space"), "6")
    top.set(qn("w:color"), color_hex)
    pBdr.append(top)
    pPr.append(pBdr)


def _add_bottom_border(paragraph, color_hex: str, size: int = 6):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_page_number_field(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def _add_section_heading(doc: Document, num: str, title: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    num_run = p.add_run(f"{num}   ")
    num_run.font.size = Pt(10)
    num_run.font.bold = True
    num_run.font.color.rgb = COLOR_BRASS
    title_run = p.add_run(title)
    title_run.font.size = Pt(16)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_INK
    _add_bottom_border(p, "E2E8F0")
    return p


def _risk_level(findings: List[Dict]) -> str:
    bloquant = sum(1 for f in findings if f.get("criticite") == "BLOQUANT")
    majeur = sum(1 for f in findings if f.get("criticite") == "MAJEUR")
    if bloquant > 0:
        return "CRITIQUE"
    if majeur >= 3:
        return "ÉLEVÉ"
    if majeur > 0:
        return "MODÉRÉ"
    if findings:
        return "FAIBLE"
    return "AUCUN"


def _make_criticality_chart(bloquant: int, majeur: int, mineur: int) -> io.BytesIO:
    fig, ax = plt.subplots(figsize=(5.6, 2.4), dpi=200)
    labels = ["Bloquants", "Majeurs", "Mineurs"]
    values = [bloquant, majeur, mineur]
    colors = ["#DC2626", "#D97706", "#64748B"]
    bars = ax.barh(labels, values, color=colors, height=0.55)
    ax.invert_yaxis()
    for spine in ("top", "right", "left"):
        ax.spines[spine].set_visible(False)
    ax.tick_params(left=False, labelsize=10)
    ax.set_xticks([])
    max_v = max(values + [1])
    for bar, v in zip(bars, values):
        ax.text(bar.get_width() + max_v * 0.02, bar.get_y() + bar.get_height() / 2,
                 str(v), va="center", fontsize=10, color="#0B1220")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", transparent=True)
    plt.close(fig)
    buf.seek(0)
    return buf


def _make_business_impact_chart(temps_manuel_min: float, temps_auto_min: float) -> io.BytesIO:
    fig, ax = plt.subplots(figsize=(5.6, 1.9), dpi=200)
    labels = ["Audit manuel (estimé)", "Audit automatisé (mesuré)"]
    values = [temps_manuel_min, temps_auto_min]
    colors = ["#94A3B8", "#2563EB"]
    bars = ax.barh(labels, values, color=colors, height=0.5)
    ax.invert_yaxis()
    for spine in ("top", "right", "left", "bottom"):
        ax.spines[spine].set_visible(False)
    ax.tick_params(left=False, labelsize=9.5)
    ax.set_xticks([])
    max_v = max(values + [1])
    for bar, v in zip(bars, values):
        label = f"{v:.0f} min" if v < 60 else f"{v / 60:.1f} h"
        ax.text(bar.get_width() + max_v * 0.02, bar.get_y() + bar.get_height() / 2,
                 label, va="center", fontsize=9.5, color="#0B1220")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", transparent=True)
    plt.close(fig)
    buf.seek(0)
    return buf


# ============================================================
# GÉNÉRATION DU RAPPORT
# ============================================================
def generate_word_report(
    reference: str,
    findings: List[Dict],
    coverage_rate: Optional[float] = None,
    health_score: Optional[int] = None,
    health_niveau: Optional[str] = None,
    logo_path: Optional[Path] = None,
    date_str: Optional[str] = None,
    business_impact: Optional[Dict] = None,
    qa_history: Optional[List[Dict]] = None,
) -> bytes:
    """Construit le rapport d'audit Word (charte Talan) et retourne les bytes du .docx.

    business_impact (optionnel) : dict avec les clés 'temps_auto_min',
    'temps_manuel_min', 'gain_min', 'gain_pct', 'cout_gagne', 'taux_horaire',
    'n_kpis' — si fourni, ajoute une section 03 dédiée en annexe du rapport.

    qa_history (optionnel) : liste de {'question': str, 'answer': str} — si
    fournie et non vide, ajoute une annexe listant les échanges avec
    l'assistant Q&A de la session en cours.
    """
    date_str = date_str or datetime.date.today().strftime("%d %B %Y")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = COLOR_INK

    # ---------- PAGE DE GARDE ----------
    if logo_path and Path(logo_path).exists():
        p_logo = doc.add_paragraph()
        p_logo.add_run().add_picture(str(logo_path), height=Cm(1.3))

    band_p = doc.add_paragraph()
    _add_top_border(band_p, "1E3A5F", size=24)

    kicker = doc.add_paragraph()
    kicker_run = kicker.add_run("AUDIT DE MIGRATION · DOCUMENT CONFIDENTIEL · TALAN")
    kicker_run.font.size = Pt(9)
    kicker_run.font.bold = True
    kicker_run.font.color.rgb = COLOR_INK_SOFT

    title = doc.add_paragraph()
    title_run = title.add_run("Qlik Sense vers Power BI")
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_INK

    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run("Rapport d'audit de migration — synthèse et constats détaillés")
    subtitle_run.font.size = Pt(12)
    subtitle_run.font.color.rgb = COLOR_INK_SOFT

    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(8)
    r = meta.add_run("Référence : ")
    r.font.bold = True
    meta.add_run(f"{reference}      ")
    r = meta.add_run("Date : ")
    r.font.bold = True
    meta.add_run(date_str)

    doc.add_paragraph()

    # ---------- 01. RÉSUMÉ EXÉCUTIF ----------
    _add_section_heading(doc, "01", "Résumé exécutif")

    bloquant = sum(1 for f in findings if f.get("criticite") == "BLOQUANT")
    majeur = sum(1 for f in findings if f.get("criticite") == "MAJEUR")
    mineur = sum(1 for f in findings if f.get("criticite") == "MINEUR")
    niveau = health_niveau or _risk_level(findings)

    summary_p = doc.add_paragraph()
    summary_p.add_run(
        f"L'audit automatisé a identifié {len(findings)} point(s) à examiner "
        f"({bloquant} bloquant(s), {majeur} majeur(s), {mineur} mineur(s)). "
        f"Le niveau de risque global de cette migration est évalué à "
    )
    risk_run = summary_p.add_run(niveau)
    risk_run.font.bold = True
    risk_run.font.color.rgb = NIVEAU_COLORS.get(niveau, COLOR_ACCENT)
    summary_p.add_run(".")

    if bloquant > 0:
        blocking_p = doc.add_paragraph()
        blocking_run = blocking_p.add_run(
            "Les points bloquants nécessitent une validation manuelle par un "
            "consultant avant de considérer la migration comme terminée."
        )
        blocking_run.font.italic = True
        blocking_run.font.color.rgb = COLOR_DANGER

    # ---------- Stats + graphique ----------
    doc.add_paragraph()
    stats_table = doc.add_table(rows=2, cols=4)
    stats_table.autofit = False
    headers = ["Bloquants", "Majeurs", "Mineurs", "Score de santé"]
    values = [str(bloquant), str(majeur), str(mineur),
              f"{health_score}/100" if health_score is not None else "—"]
    colors_hdr = ["FEE2E2", "FEF3C7", "F1F5F9", "DBEAFE"]
    text_colors = [COLOR_DANGER, COLOR_WARN, COLOR_NEUTRAL, COLOR_ACCENT2]
    col_width = Cm(4.15)
    for i in range(4):
        for row_idx in (0, 1):
            stats_table.rows[row_idx].cells[i].width = col_width
        _shade_cell(stats_table.rows[0].cells[i], colors_hdr[i])
        _set_cell_text(stats_table.rows[0].cells[i], headers[i], color=COLOR_INK, bold=True,
                        size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(stats_table.rows[1].cells[i], values[i], color=text_colors[i], bold=True,
                        size=17, align=WD_ALIGN_PARAGRAPH.CENTER)

    if bloquant or majeur or mineur:
        doc.add_paragraph()
        chart_buf = _make_criticality_chart(bloquant, majeur, mineur)
        doc.add_picture(chart_buf, width=Cm(14.5))

    if coverage_rate is not None:
        cov_p = doc.add_paragraph()
        cov_p.paragraph_format.space_before = Pt(6)
        r = cov_p.add_run("Taux de couverture fonctionnelle (Module B) : ")
        r.font.bold = True
        cov_p.add_run(f"{coverage_rate:.0f}%")

    # ---------- 02. DÉTAIL DES CONSTATS ----------
    doc.add_page_break()
    _add_section_heading(doc, "02", "Détail des constats")

    ordre = {"BLOQUANT": 0, "MAJEUR": 1, "MINEUR": 2}
    findings_sorted = sorted(findings, key=lambda x: ordre.get(x.get("criticite"), 3))

    if not findings_sorted:
        doc.add_paragraph("Aucun constat à signaler.")
    else:
        for f in findings_sorted:
            crit = f.get("criticite", "MINEUR")
            crit_color, crit_bg = CRIT_STYLE.get(crit, CRIT_STYLE["MINEUR"])

            tbl = doc.add_table(rows=1, cols=2)
            tbl.autofit = False
            tag_cell, title_cell = tbl.rows[0].cells
            tag_cell.width = Cm(2.4)
            title_cell.width = Cm(14.2)
            tag_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            title_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _shade_cell(tag_cell, crit_bg)
            _set_cell_text(tag_cell, crit, color=crit_color, bold=True, size=8,
                            align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_cell_text(title_cell, f.get("libelle", ""), color=COLOR_INK, bold=True, size=11)

            origin_p = doc.add_paragraph()
            origin_p.paragraph_format.space_before = Pt(4)
            origin_p.add_run("Origine : ").font.bold = True
            origin_p.add_run(f.get("source_module", ""))

            if f.get("detail"):
                p = doc.add_paragraph()
                p.add_run("Détail : ").font.bold = True
                p.add_run(f.get("detail", ""))

            if f.get("diagnostic"):
                p = doc.add_paragraph()
                p.add_run("Diagnostic : ").font.bold = True
                p.add_run(f.get("diagnostic", ""))

            if f.get("recommandation"):
                p = doc.add_paragraph()
                p.add_run("Recommandation : ").font.bold = True
                rec_run = p.add_run(f.get("recommandation", ""))
                rec_run.font.color.rgb = COLOR_OK

            doc.add_paragraph()  # spacer

    # ---------- 03. IMPACT BUSINESS (optionnel) ----------
    if business_impact:
        doc.add_page_break()
        _add_section_heading(doc, "03", "Impact business")

        intro_p = doc.add_paragraph()
        intro_p.add_run(
            "Temps réellement mesuré par l'outil, comparé à une estimation du temps "
            "qu'un audit manuel équivalent aurait pris à un consultant."
        )

        n_kpis = business_impact.get("n_kpis", 0)
        temps_auto_min = business_impact.get("temps_auto_min", 0.0)
        temps_manuel_min = business_impact.get("temps_manuel_min", 0.0)
        gain_min = business_impact.get("gain_min", 0.0)
        gain_pct = business_impact.get("gain_pct", 0.0)
        cout_gagne = business_impact.get("cout_gagne", 0.0)
        taux_horaire = business_impact.get("taux_horaire")

        def _fmt_duree(minutes: float) -> str:
            if minutes < 1:
                return f"{minutes * 60:.0f} sec"
            if minutes < 60:
                return f"{minutes:.1f} min"
            h = int(minutes // 60)
            m = int(minutes % 60)
            return f"{h}h{m:02d}"

        doc.add_paragraph()
        bi_table = doc.add_table(rows=2, cols=4)
        bi_table.autofit = False
        bi_headers = ["Temps automatisé", "Temps manuel estimé", "Temps gagné", "Réduction"]
        bi_values = [_fmt_duree(temps_auto_min), _fmt_duree(temps_manuel_min),
                     _fmt_duree(gain_min), f"{gain_pct:.0f}%"]
        bi_colors_hdr = ["D1FAE5", "F1F5F9", "D1FAE5", "DBEAFE"]
        bi_text_colors = [COLOR_OK, COLOR_NEUTRAL, COLOR_OK, COLOR_ACCENT2]
        for i in range(4):
            for row_idx in (0, 1):
                bi_table.rows[row_idx].cells[i].width = col_width
            _shade_cell(bi_table.rows[0].cells[i], bi_colors_hdr[i])
            _set_cell_text(bi_table.rows[0].cells[i], bi_headers[i], color=COLOR_INK, bold=True,
                            size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_cell_text(bi_table.rows[1].cells[i], bi_values[i], color=bi_text_colors[i], bold=True,
                            size=15, align=WD_ALIGN_PARAGRAPH.CENTER)

        doc.add_paragraph()
        chart_buf = _make_business_impact_chart(temps_manuel_min, temps_auto_min)
        doc.add_picture(chart_buf, width=Cm(14.5))

        if cout_gagne and taux_horaire:
            cost_p = doc.add_paragraph()
            cost_p.paragraph_format.space_before = Pt(10)
            r = cost_p.add_run("Économie estimée sur ce run : ")
            r.font.bold = True
            amount_run = cost_p.add_run(f"{cout_gagne:,.0f} €".replace(",", " "))
            amount_run.font.bold = True
            amount_run.font.color.rgb = COLOR_ACCENT2
            cost_p.add_run(f" (base : {taux_horaire:.0f} €/heure).")

        basis_p = doc.add_paragraph()
        basis_p.paragraph_format.space_before = Pt(4)
        basis_run = basis_p.add_run(
            f"Base de calcul : {n_kpis} KPI(s) réconcilié(s) automatiquement, "
            f"{len(findings)} constat(s) diagnostiqué(s) au total."
        )
        basis_run.font.size = Pt(9)
        basis_run.font.color.rgb = COLOR_INK_SOFT

        note_p = doc.add_paragraph()
        note_p.paragraph_format.space_before = Pt(10)
        note_run = note_p.add_run(
            "Note de méthode : le temps automatisé est mesuré directement (chronométré "
            "autour de l'extraction et de l'analyse de couverture), pas estimé. Le temps "
            "manuel est une hypothèse raisonnable, à ajuster selon le contexte réel "
            "(complexité du modèle, expérience du consultant), pas une mesure de terrain."
        )
        note_run.font.italic = True
        note_run.font.size = Pt(9)
        note_run.font.color.rgb = COLOR_INK_SOFT

    # ---------- 04. ANNEXE — ASSISTANT Q&A (optionnel) ----------
    if qa_history:
        doc.add_page_break()
        section_num = "04" if business_impact else "03"
        _add_section_heading(doc, section_num, "Annexe — Échanges avec l'assistant Q&A")

        intro_qa = doc.add_paragraph()
        intro_qa.add_run(
            "Questions posées à l'assistant pendant la session d'audit et réponses "
            "obtenues, à partir des constats générés — fourni à titre indicatif, "
            "à valider comme le reste des éléments assistés par IA de ce rapport."
        )

        for i, turn in enumerate(qa_history, 1):
            q_p = doc.add_paragraph()
            q_p.paragraph_format.space_before = Pt(10)
            q_label = q_p.add_run(f"Q{i}. ")
            q_label.font.bold = True
            q_label.font.color.rgb = COLOR_ACCENT2
            q_p.add_run(turn.get("question", "")).font.bold = True

            a_p = doc.add_paragraph()
            a_p.paragraph_format.space_before = Pt(2)
            a_label = a_p.add_run("R. ")
            a_label.font.bold = True
            a_label.font.color.rgb = COLOR_INK_SOFT
            a_p.add_run(turn.get("answer", ""))

    # ---------- PIED DE PAGE ----------
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run(
        f"Audit de migration Qlik Sense vers Power BI · Talan · "
        f"{datetime.date.today().year} · Document confidentiel · Page "
    )
    run.font.size = Pt(8)
    run.font.color.rgb = COLOR_INK_SOFT
    _add_page_number_field(footer_p)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
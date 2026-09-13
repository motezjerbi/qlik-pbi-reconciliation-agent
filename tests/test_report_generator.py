"""
Tests du générateur de rapport Word (src/orchestrator/word_report_generator.py).

Vérifie que le document produit est un .docx valide et exploitable (rouvert
avec python-docx), pas seulement que la fonction ne plante pas.
"""
import io

from docx import Document

import orchestrator.word_report_generator as wrg


def _findings():
    return [
        {"criticite": "BLOQUANT", "libelle": "Current Year Sales — ECART_VALEUR",
         "source_module": "Réconciliation de données", "detail": "Qlik=0 | PBI=4100",
         "diagnostic": "Filtre d'année différent.", "recommandation": "Aligner Max(Year).", "statut": "ECART_VALEUR"},
        {"criticite": "MAJEUR", "libelle": "Colonne malformée",
         "source_module": "Module A - Structural Gaps", "detail": "ChannelID;ChannelName",
         "diagnostic": "Erreur Power Query.", "recommandation": "Corriger le délimiteur.", "statut": ""},
    ]


class TestRiskLevel:
    def test_any_bloquant_is_critique(self):
        assert wrg._risk_level([{"criticite": "BLOQUANT"}, {"criticite": "MINEUR"}]) == "CRITIQUE"

    def test_three_or_more_majeur_is_eleve(self):
        assert wrg._risk_level([{"criticite": "MAJEUR"}] * 3) == "ÉLEVÉ"

    def test_one_or_two_majeur_is_modere(self):
        assert wrg._risk_level([{"criticite": "MAJEUR"}]) == "MODÉRÉ"

    def test_only_mineur_is_faible(self):
        assert wrg._risk_level([{"criticite": "MINEUR"}]) == "FAIBLE"

    def test_no_findings_is_aucun(self):
        assert wrg._risk_level([]) == "AUCUN"


class TestGenerateWordReport:
    def test_produces_non_empty_bytes(self):
        data = wrg.generate_word_report(reference="QA-TEST", findings=_findings(), health_score=25, health_niveau="CRITIQUE")
        assert isinstance(data, bytes)
        assert len(data) > 1000  # un .docx minimal fait déjà plusieurs Ko (structure ZIP/XML)

    def test_output_is_a_valid_docx(self):
        data = wrg.generate_word_report(reference="QA-TEST", findings=_findings())
        doc = Document(io.BytesIO(data))  # doit s'ouvrir sans lever d'exception
        assert len(doc.paragraphs) > 0

    def test_findings_content_appears_in_document(self):
        data = wrg.generate_word_report(reference="QA-TEST", findings=_findings())
        doc = Document(io.BytesIO(data))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                full_text += "\n" + "\n".join(cell.text for cell in row.cells)
        assert "Current Year Sales" in full_text
        assert "Colonne malformée" in full_text

    def test_empty_findings_does_not_crash(self):
        data = wrg.generate_word_report(reference="QA-TEST", findings=[])
        doc = Document(io.BytesIO(data))
        assert len(doc.paragraphs) > 0

    def test_reference_appears_in_document(self):
        data = wrg.generate_word_report(reference="QA-UNIQUE-REF-123", findings=_findings())
        doc = Document(io.BytesIO(data))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "QA-UNIQUE-REF-123" in full_text